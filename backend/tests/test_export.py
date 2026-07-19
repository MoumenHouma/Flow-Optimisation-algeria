"""Route export (F5) — PDF + Excel byte generation. No DB required."""

import io
import uuid
from datetime import time
from types import SimpleNamespace

from openpyxl import load_workbook

from routeopt.modules.routes.export import _is_arabic, _shape, build_excel, build_pdf


def _fixture():
    d1 = SimpleNamespace(
        id=uuid.uuid4(),
        address="12 Rue Didouche Mourad, Alger",
        order_id="CMD-1",
        time_window_start=time(9, 0),
        time_window_end=time(12, 0),
        weight=5.2,
        customer_phone="0550123456",
    )
    d2 = SimpleNamespace(
        id=uuid.uuid4(),
        address="45 Bd Mohamed V",
        order_id=None,
        time_window_start=None,
        time_window_end=None,
        weight=8,
        customer_phone=None,
    )
    stops = [
        SimpleNamespace(delivery_id=d2.id, sequence=1),  # deliberately out of order
        SimpleNamespace(delivery_id=d1.id, sequence=0),
    ]
    route = SimpleNamespace(id=uuid.uuid4(), total_distance_m=12345, stops=stops)
    deliveries = {d1.id: d1, d2.id: d2}
    return route, deliveries


def test_build_excel_orders_stops_and_includes_data():
    route, deliveries = _fixture()
    content = build_excel(route, deliveries)
    assert content[:2] == b"PK"  # xlsx is a zip

    wb = load_workbook(io.BytesIO(content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    header_idx = next(i for i, r in enumerate(rows) if r and r[0] == "#")
    body = rows[header_idx + 1 :]
    # Ordered by sequence: stop 1 is the first delivery
    assert body[0][0] == "1"
    assert body[0][1] == "12 Rue Didouche Mourad, Alger"
    assert body[0][3] == "09:00–12:00"
    assert body[1][0] == "2"
    assert body[1][1] == "45 Bd Mohamed V"


def test_build_pdf_returns_pdf_bytes():
    route, deliveries = _fixture()
    content = build_pdf(route, deliveries)
    assert content[:5] == b"%PDF-"
    assert len(content) > 800  # non-trivial document


def test_is_arabic_detection():
    assert _is_arabic("12 شارع ديدوش مراد")
    assert not _is_arabic("12 Rue Didouche Mourad")
    assert not _is_arabic("")


def test_shape_reorders_arabic_and_leaves_latin():
    latin = "12 Rue Didouche"
    assert _shape(latin) == latin  # untouched

    arabic = "شارع"
    shaped = _shape(arabic)
    assert shaped != arabic  # reshaped to presentation forms + bidi-reordered
    # Reshaping maps to Arabic Presentation Forms-B (U+FE70..U+FEFF)
    assert any("ﹰ" <= c <= "﻿" for c in shaped)


def test_build_pdf_with_arabic_address_builds():
    d = SimpleNamespace(
        id=uuid.uuid4(),
        address="12 شارع ديدوش مراد, الجزائر",
        order_id="CMD-9",
        time_window_start=None,
        time_window_end=None,
        weight=3,
        customer_phone="0550000000",
    )
    stops = [SimpleNamespace(delivery_id=d.id, sequence=0)]
    route = SimpleNamespace(id=uuid.uuid4(), total_distance_m=1000, stops=stops)
    content = build_pdf(route, {d.id: d})
    assert content[:5] == b"%PDF-"
